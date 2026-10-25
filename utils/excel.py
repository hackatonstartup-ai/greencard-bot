import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path
from typing import List, Dict, Any

def export_applications_to_excel(applications: List[Dict[str, Any]], output_path: Path) -> Path:
    wb = openpyxl.Workbook()

    # Sheet 1: Barcha Asosiy Anketalar
    ws = wb.active
    ws.title = "Anketalar (Asosiy)"

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = [
        "ID", "Status", "Confirmation Code", "Ariza Sanasi",
        "Pasport rasm fayli", "Pasport Seriya", "Pasport Muddati",
        "Familiya (Last Name)", "Ism (First Name)",
        "Jinsi", "Tug'ilgan sana", "Tug'ilgan viloyat", "Tug'ilgan shahar/tuman", "Tug'ilgan davlat",
        "Rasm fayli",
        "Yashash davlati", "Yashash viloyati", "Yashash tumani", "Manzil (Ko'cha/Uy)",
        "Telefon raqam", "Email",
        "Ta'lim darajasi", "Oilaviy holati", "Farzandlar soni",
        "Turmush o'rtog'i ma'lumoti", "Farzandlar ma'lumoti"
    ]

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for app in applications:
        spouse_info = "Yo'q"
        if app.get("spouse"):
            sp = app["spouse"]
            spouse_info = (
                f"{sp.get('last_name', '')} {sp.get('first_name', '')} | "
                f"Tug'ilgan: {sp.get('birth_date', '')} | {sp.get('gender', '')} | "
                f"{sp.get('birth_region', '')} {sp.get('birth_city', '')}, {sp.get('birth_country', '')} | "
                f"Ta'lim: {sp.get('education_level', '')} | "
                f"Rasm: {sp.get('photo_path', '')}"
            )

        children_info = "Yo'q"
        if app.get("children"):
            ch_list = []
            for ch in app["children"]:
                ch_list.append(
                    f"#{ch.get('child_index')}: {ch.get('last_name', '')} {ch.get('first_name', '')} | "
                    f"{ch.get('birth_date', '')} | {ch.get('gender', '')} | "
                    f"{ch.get('birth_region', '')} {ch.get('birth_city', '')} | Rasm: {ch.get('photo_path', '')}"
                )
            children_info = "\n".join(ch_list)

        row_data = [
            app.get("id"),
            app.get("status", "pending"),
            app.get("confirmation_number") or "Kiritilmagan",
            app.get("created_at"),
            app.get("passport_photo_path") or "",
            app.get("passport_number") or "",
            app.get("passport_expiry") or "",
            app.get("last_name"),
            app.get("first_name"),
            app.get("gender"),
            app.get("birth_date"),
            app.get("birth_region") or "",
            app.get("birth_city") or "",
            app.get("birth_country"),
            app.get("photo_path") or "",
            app.get("live_country") or app.get("country"),
            app.get("district_state") or "",
            app.get("city") or "",
            app.get("address_line1"),
            app.get("phone_number"),
            app.get("email") or "",
            app.get("education_level"),
            app.get("marital_status"),
            app.get("children_count", 0),
            spouse_info,
            children_info
        ]
        ws.append(row_data)

    # Sheet 2: Turmush o'rtoqlari
    ws_spouse = wb.create_sheet(title="Turmush o'rtoqlari")
    spouse_headers = [
        "Asosiy Anketa ID", "Familiya", "Ism",
        "Jinsi", "Tug'ilgan sana", "Tug'ilgan viloyat", "Tug'ilgan shahar/tuman", "Tug'ilgan davlat", "Ta'lim", "Rasm yo'li"
    ]
    ws_spouse.append(spouse_headers)
    for col_num in range(1, len(spouse_headers) + 1):
        cell = ws_spouse.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = PatternFill(start_color="333f48", end_color="333f48", fill_type="solid")
        cell.alignment = center_align

    for app in applications:
        if app.get("spouse"):
            sp = app["spouse"]
            ws_spouse.append([
                app.get("id"),
                sp.get("last_name"),
                sp.get("first_name"),
                sp.get("gender"),
                sp.get("birth_date"),
                sp.get("birth_region") or "",
                sp.get("birth_city") or "",
                sp.get("birth_country"),
                sp.get("education_level") or "",
                sp.get("photo_path")
            ])

    # Sheet 3: Farzandlar
    ws_children = wb.create_sheet(title="Farzandlar")
    child_headers = [
        "Asosiy Anketa ID", "Farzand tartibi", "Familiya", "Ism",
        "Jinsi", "Tug'ilgan sana", "Tug'ilgan viloyat", "Tug'ilgan shahar/tuman", "Tug'ilgan davlat", "Rasm yo'li"
    ]
    ws_children.append(child_headers)
    for col_num in range(1, len(child_headers) + 1):
        cell = ws_children.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = PatternFill(start_color="2e5b32", end_color="2e5b32", fill_type="solid")
        cell.alignment = center_align

    for app in applications:
        for ch in app.get("children", []):
            ws_children.append([
                app.get("id"),
                ch.get("child_index"),
                ch.get("last_name"),
                ch.get("first_name"),
                ch.get("gender"),
                ch.get("birth_date"),
                ch.get("birth_region") or "",
                ch.get("birth_city") or "",
                ch.get("birth_country"),
                ch.get("photo_path")
            ])

    # Column widths
    for sheet in [ws, ws_spouse, ws_children]:
        for col in sheet.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val_str = str(cell.value or "")
                if "\n" in val_str:
                    lines = val_str.split("\n")
                    max_len = max(max_len, max(len(l) for l in lines))
                else:
                    max_len = max(max_len, len(val_str))
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
